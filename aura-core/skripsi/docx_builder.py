import re
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_styled_document():
    doc = Document()
    
    # Configure A4 Margins: Top 4cm, Left 4cm, Bottom 3cm, Right 3cm (Standar Resmi UII)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.0)
    section.left_margin = Cm(4.0)
    section.bottom_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    
    # Configure Normal Style (Times New Roman 12pt, 1.5 line spacing)
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    return doc

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_academic_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_clean_paragraph(doc, text, style='Normal', align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, space_before=0, space_after=6, line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    # Process bold and italic inside text
    # Regex to split on **bold** or *italic*
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.italic = True
        else:
            run = p.add_run(token)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = bold
            run.italic = italic
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_math_formula(doc, formula_text, eq_num=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    
    clean_formula = formula_text.strip().replace('$$', '').replace('$', '').strip()
    run = p.add_run(clean_formula)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.italic = True
    
    if eq_num:
        run_num = p.add_run(f"\t\t\t\t({eq_num})")
        run_num.font.name = 'Times New Roman'
        run_num.font.size = Pt(11)
        run_num.bold = True
    return p

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_academic_table_borders(table)
    
    # Format Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title.strip()
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_shading(hdr_cells[i], "F0F4F8") # Subtle elegant academic blue-gray
        hdr_p = hdr_cells[i].paragraphs[0]
        hdr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_p.paragraph_format.space_before = Pt(0)
        hdr_p.paragraph_format.space_after = Pt(0)
        hdr_p.paragraph_format.line_spacing = 1.0
        for r in hdr_p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(0, 32, 96)
            
    # Format Rows
    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "FFFFFF" if row_idx % 2 == 0 else "FBFBFC"
        for col_idx, cell_value in enumerate(row_data):
            if col_idx < len(row_cells):
                cell = row_cells[col_idx]
                cell.text = cell_value.strip()
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                set_cell_shading(cell, bg_color)
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.space_before = Pt(0)
                cell_p.paragraph_format.space_after = Pt(0)
                cell_p.paragraph_format.line_spacing = 1.15
                
                # Align numbers/short text to center, long text to left
                val_clean = cell_value.strip()
                if len(val_clean) <= 12 or re.match(r'^[0-9.,%+\-s]+$', val_clean) or val_clean.startswith('S') or val_clean.startswith('CS'):
                    cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                for r in cell_p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(9.5)
                    if "**" in val_clean or val_clean.isupper() and len(val_clean) < 15:
                        r.bold = True
                        
    # Set explicit widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_callout_box(doc, text_content, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.rows[0].cells[0]
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_shading(cell, "F4F6F9")
    
    # Add left border accent (Blue #003366)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    
    if title:
        run_title = p.add_run(f"[{title}]\n")
        run_title.font.name = 'Times New Roman'
        run_title.font.size = Pt(10.5)
        run_title.bold = True
        run_title.font.color.rgb = RGBColor(0, 51, 102)
        
    run_txt = p.add_run(text_content.strip())
    run_txt.font.name = 'Times New Roman'
    run_txt.font.size = Pt(10.5)
    run_txt.italic = True
    run_txt.font.color.rgb = RGBColor(30, 30, 30)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_flow_diagram_table(doc):
    # Professional replacement for raw mermaid flowchart
    headers = ["Fase", "Tahapan Penelitian", "Aktivitas Kunci & Luaran"]
    rows = [
        ["Fase 1", "Identifikasi Masalah & Studi Literatur", "Menelaah kelemahan konsultasi akademik, risiko halusinasi RAG tahap tunggal, dan teori RAG Triad."],
        ["Fase 2", "Akuisisi Korpus & Semantic Chunking", "Pengumpulan dokumen resmi FTI UII (Pedoman UII, Skripsi FTI), kurasi, dan segmentasi 500-800 token."],
        ["Fase 3", "Perancangan Arsitektur Pure-Go Two-Stage RAG", "Konstruksi orkestrator Go native, Dense ANN (Top-20), Neural Reranker (Top-8), dan In-Memory Parser KHS/KRS."],
        ["Fase 4", "Pengembangan Benchmark UII-Bench-50 & Judge", "Pembuatan dataset 50 skenario (5 klaster) dan penguji otomatis LLM-as-Judge (cmd/bench)."],
        ["Fase 5", "Pengujian Kuantitatif, Studi Ablasi & Uji Kasus", "Evaluasi 50 skenario, pengujian ablasi (+10.58%), profiling latensi komputasi, dan 15 studi kasus riil."],
        ["Fase 6", "Analisis Pembahasan & Kesimpulan", "Sintesis hasil evaluasi RAG Triad, evaluasi kasus batas (edge cases), perumusan kesimpulan dan saran."]
    ]
    col_widths = [Cm(1.8), Cm(5.2), Cm(7.0)]
    add_styled_table(doc, headers, rows, col_widths)

print("Helper definitions loaded successfully.")
EOF
